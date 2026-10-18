"""Application ingestion checks for common spreadsheet/CSV export variants."""
import io
import pandas as pd
import pytest


def payloads():
    base = pd.DataFrame({'name':['José','München','Zoë'], 'value':[1.25,0.,-3.5], 'target':[1,0,1]})
    yield 'utf8.csv', base.to_csv(index=False).encode('utf-8'), base
    yield 'bom.csv', base.to_csv(index=False).encode('utf-8-sig'), base
    yield 'windows.csv', base.to_csv(index=False).encode('cp1252'), base
    yield 'unicode.csv', base.to_csv(index=False).encode('utf-16'), base
    for separator,name in [(';','semicolon'), ('\t','tab'), ('|','pipe')]:
        yield name+'.csv', base.to_csv(index=False,sep=separator).encode(), base
    yield 'decimal-comma.csv', base.to_csv(index=False,sep=';',decimal=',').encode(), base
    quoted = pd.DataFrame({'text':['comma,inside','first\nsecond','say "hello"'], 'value':[1,2,3]})
    yield 'quoted.csv', quoted.to_csv(index=False).encode(), quoted
    excel = io.BytesIO(); base.to_excel(excel,index=False)
    yield 'workbook.xlsx', excel.getvalue(), base
    from xlwt import Workbook
    book=Workbook(); sheet=book.add_sheet('Data')
    for c,name in enumerate(base.columns): sheet.write(0,c,name)
    for r,row in enumerate(base.itertuples(index=False,name=None),1):
        for c,value in enumerate(row): sheet.write(r,c,value)
    legacy=io.BytesIO(); book.save(legacy)
    yield 'legacy.xls', legacy.getvalue(), base


EXPORTS = list(payloads())


@pytest.mark.parametrize('name,content,expected', EXPORTS, ids=[case[0] for case in EXPORTS])
def test_supported_export_variants_preserve_values(client,accounts,name,content,expected):
    response=client.post('/api/profile',headers=accounts[0],files={'file':(name,content)})
    assert response.status_code==200,response.text
    profile=response.json(); pid=profile['project_id']
    assert profile['overview']['total_rows']==len(expected)
    assert profile['column_names']==list(expected.columns)
    downloaded=client.get(f'/api/datasets/{pid}/download',headers=accounts[0])
    assert downloaded.status_code==200
    restored=pd.read_csv(io.BytesIO(downloaded.content))
    pd.testing.assert_frame_equal(restored,expected,check_dtype=False)


@pytest.mark.parametrize('name,content,expected', EXPORTS, ids=[case[0] for case in EXPORTS])
def test_direct_pipeline_loader_matches_api_exports(tmp_path,name,content,expected):
    from backend.agents.csv_loader import load_file_dataframe
    pd.testing.assert_frame_equal(load_file_dataframe(content,name),expected,check_dtype=False)
    path=tmp_path/name;path.write_bytes(content)
    pd.testing.assert_frame_equal(load_file_dataframe(str(path),name),expected,check_dtype=False)


INVALID_FILES = [
    ('empty.csv',b''), ('header.csv',b'a,b\n'), ('nulls.csv',b'a,b\n\x00,\x01\n'),
    ('ragged.csv',b'a,b\n1,2,3\n4,5,6\n'), ('truncated.csv',b'a,b\n1,2\n3\n'),
    ('broken.csv',b'a,b\n"unterminated,2\n'), ('corrupt.xlsx',b'not an Excel file'),
    ('corrupt.xls',b'not an Excel file'), ('data.json',b'[{"a":1}]'),
    ('data.parquet',b'PAR1'), ('data.zip',b'PK\x03\x04'),
]


@pytest.mark.parametrize('name,content', INVALID_FILES, ids=[case[0] for case in INVALID_FILES])
def test_invalid_files_are_rejected_without_server_errors(client,accounts,name,content):
    response=client.post('/api/profile',headers=accounts[0],files={'file':(name,content)})
    assert response.status_code==400,response.text
    assert 'detail' in response.json()


def test_excel_dates_blank_columns_and_duplicate_headers(client,accounts):
    from openpyxl import Workbook
    book=Workbook(); sheet=book.active
    sheet.append([' date ', 'amount', 'amount', 'amount_2', None])
    for n in range(30): sheet.append([(pd.Timestamp('2025-01-01')+pd.Timedelta(days=n)).to_pydatetime(),n,n*2,n*3,None])
    sheet2=book.create_sheet('Other'); sheet2.append(['ignored']); sheet2.append(['second sheet'])
    data=io.BytesIO();book.save(data)
    response=client.post('/api/profile',headers=accounts[0],files={'file':('dates.xlsx',data.getvalue())})
    assert response.status_code==200,response.text
    result=response.json()
    assert result['overview']['total_rows']==30
    assert len(result['column_names'])==len(set(result['column_names']))
    assert result['col_types']['date']=='datetime'


def test_infinite_numeric_csv_can_be_profiled_and_cleaned(client,accounts):
    content=b'feature,target\n1,0\ninf,1\n-inf,0\n,1\n5,0\n'
    response=client.post('/api/profile',headers=accounts[0],files={'file':('infinite.csv',content)})
    assert response.status_code==200,response.text
    pid=response.json()['project_id']
    cleaned=client.post('/api/clean',headers=accounts[0],json={'project_id':pid,'target_column':'target'})
    assert cleaned.status_code==200,cleaned.text
    frame=pd.DataFrame(cleaned.json()['preview'])
    assert frame['feature'].tolist()==[1,3,3,3,5]
    assert frame['target'].tolist()==[0,1,0,1,0]
    assert cleaned.json()['audit_log']['nonfinite_replacements']==[{'column':'feature','count':2}]


def test_schema_preserves_non_dates_and_parses_day_first(monkeypatch):
    import importlib
    from backend.services.profiler import profile_dataset
    schema=importlib.import_module('backend.agents.schema_agent')
    monkeypatch.setattr(schema,'invoke_llm',lambda **kwargs:'Offline deterministic fallback')
    df=pd.DataFrame({'Date':['10/03/2004','13/03/2004','01/04/2004'],
        'Time':['18.00.00','19.00.00','20.00.00'], 'updated_status':['pending','done','failed']})
    result=schema.schema_agent_node({'df_raw':df})['df_clean']
    assert result['Date'].dt.strftime('%Y-%m-%d').tolist()==['2004-03-10','2004-03-13','2004-04-01']
    pd.testing.assert_series_equal(result['Time'],df['Time'])
    pd.testing.assert_series_equal(result['updated_status'],df['updated_status'])
    assert profile_dataset(df)['col_types']['Date']=='datetime'


def test_dates_are_not_ids_and_long_text_is_not_an_identifier():
    from backend.services.profiler import profile_dataset
    df=pd.DataFrame({'Month':pd.date_range('2000-01-01',periods=144,freq='MS').strftime('%Y-%m'),
        'CustomerId':range(144),'Notes':[f'Unique customer narrative {n}' for n in range(144)]})
    types=profile_dataset(df)['col_types']
    assert types['Month']=='datetime'
    assert types['CustomerId']=='id'
    assert types['Notes']=='text'
